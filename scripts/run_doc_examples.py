"""Run the shell examples of the documentation against fixture files.

Usage:
    python scripts/run_doc_examples.py                 # every page under docs/docs
    python scripts/run_doc_examples.py docs/docs/commands/head.md
    python scripts/run_doc_examples.py --list          # print the commands without running

Every ```bash block is executed line by line (continuation lines joined) in a fresh
sandbox directory that contains the fixture files from :func:`make_fixtures`. Only lines
that call ``undatum`` (or the ``data`` alias) are run; ``cd``, ``export`` and ``mkdir`` lines
are applied too so that blocks can set up state. A block whose info string contains
``norun`` is skipped — use it for examples that need a database server, cloud storage,
an AI provider, network access, or that start a long-running server:

    ```bash norun
    undatum db query "SELECT 1" --db postgresql://user:pass@host/db
    ```

Lines with elisions (``...``, ``…``) or placeholders (``<file>``) are skipped as well.
Exit code 1 when any example fails.
"""

from __future__ import annotations

import argparse
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs" / "docs"
FENCE = re.compile(r"^\s*```(\w*)(.*)$")
# <file>, <command>, <uri>: placeholders rather than redirections.
PLACEHOLDER = re.compile(r"<[A-Za-z][\w-]*>")
SETUP_PREFIXES = ("cd ", "export ", "mkdir ")
TIMEOUT = 30
# Files named in the docs that hold the same sample rows (CSV and JSON Lines).
TABULAR_FIXTURES = (
    "sales",
    "people",
    "users",
    "orders",
    "events",
    "input",
    "public",
    "big",
    "file1",
    "file2",
    "data1",
    "data2",
    "huge",
    "clean",
    "stage",
    "other",
    "cities",
    "source",
    "local",
    "new",
    "messy",
    "api_data",
    "exclude",
    "blacklist",
    "skip",
    "current",
    "previous",
    "raw",
)
RULES = """rules:
  - field: email
    name: Email
    required: true
    type: string
    format: email
    severity: error
  - field: age
    name: Age range
    type: number
    min: 0
    max: 120
    severity: warning
  - field: status
    name: Status values
    type: string
    enum: [active, inactive, pending]
    severity: error
"""
PIPELINE = """steps:
  - name: drop_dupes
    command: dedup
    args:
      input: data.csv
      output: deduped.csv
"""


@dataclass
class Example:
    """One runnable line of a documentation code block."""

    path: pathlib.Path
    line: int
    block: int
    command: str


def collect(paths: list[pathlib.Path]) -> list[Example]:
    """Return the runnable lines of every bash block in ``paths``."""
    examples: list[Example] = []
    block_id = 0
    for path in paths:
        in_block = skip = False
        pending, start = "", 0
        for number, raw in enumerate(path.read_text(encoding="utf8").splitlines(), start=1):
            fence = FENCE.match(raw)
            if fence:
                if not in_block and fence.group(1) in ("bash", "sh", "shell"):
                    in_block, skip = True, "norun" in fence.group(2)
                    block_id += 1
                elif in_block:
                    in_block = False
                continue
            if not in_block or skip:
                continue
            line = raw.strip()
            if line.endswith("\\"):
                pending += line[:-1] + " "
                start = start or number
                continue
            command, first = (pending + line).strip(), start or number
            pending, start = "", 0
            if not command or command.startswith("#"):
                continue
            command = re.sub(r"\s+#.*$", "", command)
            if any(marker in command for marker in ("...", "…", "{{")) or PLACEHOLDER.search(
                command
            ):
                continue
            words = command.split()
            runs_cli = "undatum" in words or words[0] == "data" or "| data " in command
            if runs_cli or command.startswith(SETUP_PREFIXES):
                examples.append(Example(path, first, block_id, command))
    return examples


def make_fixtures(target: pathlib.Path) -> None:
    """Write the sample files the documentation examples refer to."""
    rows = [
        {
            "id": i,
            "user_id": 100 + i % 4,
            "name": name,
            "age": age,
            "city": city,
            "country": country,
            "email": f"{name.lower()}@example.org",
            "phone": f"+1-555-01{i:02d}",
            "amount": amount,
            "price": round(amount / 3, 2),
            "quantity": i % 5 + 1,
            "status": status,
            "category": category,
            "date": f"2024-0{i % 9 + 1}-1{i % 9}",
            "created_at": f"2024-0{i % 9 + 1}-1{i % 9}T10:00:00",
            "score": round(age * 1.5, 1),
            "tags": "a,b" if i % 2 else "c",
            "region": "north" if i % 2 else "south",
            "level": "ERROR" if i % 3 == 0 else "INFO",
            "message": f"event {i}",
        }
        for i, (name, age, city, country, amount, status, category) in enumerate(
            [
                ("Alice", 34, "Berlin", "DE", 120.5, "active", "books"),
                ("Bob", 28, "Paris", "FR", 80.0, "inactive", "games"),
                ("Carol", 45, "Berlin", "DE", 300.25, "active", "music"),
                ("Dave", 39, "Madrid", "ES", 42.0, "active", "books"),
                ("Eve", 23, "Paris", "FR", 15.75, "pending", "games"),
                ("Frank", 51, "Rome", "IT", 230.0, "active", "music"),
                ("Grace", 30, "Berlin", "DE", 99.99, "inactive", "books"),
                ("Heidi", 27, "Lisbon", "PT", 64.5, "active", "games"),
            ],
            start=1,
        )
    ]
    nested = [
        {
            "id": row["id"],
            "name": row["name"],
            "email": row["email"],
            "age": row["age"],
            "status": row["status"],
            "user": {"name": row["name"], "age": row["age"], "email": row["email"]},
            "city": {"name": row["city"], "lat": 52.5 + row["id"], "lon": 13.4},
            "address": {"city": row["city"], "country": row["country"]},
            "capital_city": {"name": row["city"], "lat": 40.0 + row["id"], "lon": 10.5},
            "items": [{"sku": f"S{row['id']}", "qty": 1}],
            "tags": ["a", "b"],
        }
        for row in rows
    ]
    _write_tabular(target, "data", rows)
    for name in TABULAR_FIXTURES:
        _write_tabular(target, name, rows)
    for name in ("nested", "nested1", "nested2", "left", "right", "skip"):
        _write_jsonl(target / f"{name}.jsonl", nested)
    _write_jsonl(target / "logs.jsonl", rows)
    (target / "data.json").write_text(json.dumps(rows, indent=1), encoding="utf8")
    (target / "quoted.csv").write_text("id,name\n1,'Smith, J'\n2,'Doe'\n", encoding="utf8")
    (target / "data.xml").write_text(
        "<items>"
        + "".join(f"<item><id>{r['id']}</id><name>{r['name']}</name></item>" for r in rows)
        + "</items>",
        encoding="utf8",
    )
    (target / "rules.yml").write_text(RULES, encoding="utf8")
    shutil.copy(target / "rules.yml", target / "validation-rules.yml")
    (target / "transform.py").write_text(
        "def process(item):\n    item['processed'] = True\n    return item\n", encoding="utf8"
    )
    (target / "query.sql").write_text("SELECT city, COUNT(*) AS n FROM data GROUP BY city\n")
    (target / "my-pipeline.yml").write_text(PIPELINE, encoding="utf8")
    (target / "api.yml").write_text("resources:\n  - name: data\n    path: data.csv\n")
    _write_derived(target, rows)


def _write_jsonl(path: pathlib.Path, rows: list[dict]) -> None:
    path.write_text("".join(json.dumps(r) + "\n" for r in rows), encoding="utf8")


def _write_tabular(target: pathlib.Path, stem: str, rows: list[dict]) -> None:
    import csv

    with open(target / f"{stem}.csv", "w", newline="", encoding="utf8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    _write_jsonl(target / f"{stem}.jsonl", rows)


def _write_derived(target: pathlib.Path, rows: list[dict]) -> None:
    """Binary and compressed fixtures written with the libraries undatum depends on."""
    import bz2
    import gzip
    import lzma
    import sqlite3
    import zipfile

    import openpyxl
    import pandas as pd
    import zstandard

    frame = pd.DataFrame(rows)
    for stem in (
        "data",
        "users",
        "people",
        "huge",
        "clean",
        "stage",
        "sales",
        "other",
        "file1",
        "file2",
        "current",
        "previous",
    ):
        frame.to_parquet(target / f"{stem}.parquet", index=False)
    csv_bytes = (target / "data.csv").read_bytes()
    jsonl_bytes = (target / "data.jsonl").read_bytes()
    (target / "data.csv.gz").write_bytes(gzip.compress(csv_bytes))
    (target / "data.jsonl.gz").write_bytes(gzip.compress(jsonl_bytes))
    (target / "data.jsonl.xz").write_bytes(lzma.compress(jsonl_bytes))
    (target / "data.jsonl.bz2").write_bytes(bz2.compress(jsonl_bytes))
    (target / "data.csv.zst").write_bytes(zstandard.ZstdCompressor().compress(csv_bytes))
    for stem in ("huge", "data", "raw"):
        (target / f"{stem}.jsonl.zst").write_bytes(zstandard.ZstdCompressor().compress(jsonl_bytes))
    (target / "events.jsonl.gz").write_bytes(gzip.compress(jsonl_bytes))
    with zipfile.ZipFile(target / "data.zip", "w") as archive:
        archive.writestr("data.jsonl", jsonl_bytes)
    for name in ("workbook", "other"):
        book = openpyxl.Workbook()
        for index, title in enumerate(("Sheet1", "Sheet2", "Cities")):
            sheet = book.active if index == 0 else book.create_sheet()
            sheet.title = title
            sheet.append(list(rows[0]))
            for row in rows:
                sheet.append(list(row.values()))
        book.save(target / f"{name}.xlsx")
    for name in ("db.db", "data.db", "database.db", "app.db", "data.sqlite"):
        conn = sqlite3.connect(target / name)
        frame.to_sql("users", conn, index=False)
        frame.to_sql("orders", conn, index=False)
        conn.close()
    (target / "data").mkdir(exist_ok=True)
    for stem in ("a", "b"):
        _write_jsonl(target / "data" / f"{stem}.jsonl", rows)


def run(examples: list[Example], verbose: bool = False) -> list[str]:
    """Run ``examples`` block by block; return failure descriptions."""
    failures: list[str] = []
    bin_dir = str(pathlib.Path(sys.executable).parent)
    env = {
        **os.environ,
        "PATH": bin_dir + os.pathsep + os.environ.get("PATH", ""),
        "NO_COLOR": "1",
        "COLUMNS": "120",
    }
    with tempfile.TemporaryDirectory(prefix="undatum-docs-") as tmp:
        fixtures = pathlib.Path(tmp) / "fixtures"
        fixtures.mkdir()
        make_fixtures(fixtures)
        current_block, sandbox, cwd = None, None, None
        for example in examples:
            if example.block != current_block:
                current_block = example.block
                sandbox = pathlib.Path(tmp) / f"block{example.block}"
                shutil.copytree(fixtures, sandbox)
                cwd = sandbox
            assert cwd is not None
            location = f"{example.path.relative_to(ROOT)}:{example.line}"
            if example.command.startswith("cd "):
                cwd = (cwd / example.command[3:].strip()).resolve()
                continue
            if example.command.startswith("export "):
                key, _, value = example.command[7:].partition("=")
                env[key.strip()] = value.strip().strip("\"'")
                continue
            try:
                result = subprocess.run(
                    ["bash", "-o", "pipefail", "-c", example.command],
                    cwd=cwd,
                    env=env,
                    capture_output=True,
                    text=True,
                    timeout=TIMEOUT,
                    check=False,
                    stdin=subprocess.DEVNULL,
                )
            except subprocess.TimeoutExpired:
                failures.append(
                    f"{location}: timed out after {TIMEOUT}s: {example.command}\n"
                    "    (a server or interactive command? mark the block `norun`)"
                )
                continue
            if verbose:
                status = "ok  " if result.returncode == 0 else "FAIL"
                print(f"{status} {location}: {example.command}")
            if result.returncode != 0:
                detail = (result.stderr or result.stdout).strip().splitlines()[-3:]
                failures.append(
                    f"{location}: exit {result.returncode}: {example.command}\n    "
                    + "\n    ".join(detail)
                )
    return failures


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("paths", nargs="*", type=pathlib.Path, help="Markdown files (default: all)")
    parser.add_argument("--list", action="store_true", help="print the examples and exit")
    parser.add_argument("-v", "--verbose", action="store_true", help="print every example")
    args = parser.parse_args(argv)
    paths = [p.resolve() for p in args.paths] or sorted(DOCS.rglob("*.md"))
    examples = collect(paths)
    if args.list:
        for example in examples:
            print(f"{example.path.relative_to(ROOT)}:{example.line}: {example.command}")
        return 0
    failures = run(examples, verbose=args.verbose)
    for failure in failures:
        print(failure)
    print(f"{len(examples) - len(failures)}/{len(examples)} documentation examples passed")
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
