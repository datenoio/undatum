"""Spreadsheets read with fastexcel: typed values, the same records on both engines."""

from __future__ import annotations

import csv
import datetime
import json
from pathlib import Path

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from undatum.common.s3_iterable import open_iterable
from undatum.core import app
from undatum.io import RowSource
from undatum.io.excel import excel_table, sheet_names

runner = CliRunner()
FIXTURES = Path(__file__).parent / "fixtures"

EXPECTED = [
    {
        "id": 1,
        "price": 1.5,
        "name": "Alice",
        "when": datetime.datetime(2020, 1, 2, 3, 4, 5),
        "day": datetime.date(2021, 5, 6),
        "flag": True,
        "code": "007",
    },
    {
        "id": 2,
        "price": 2.0,
        "name": None,
        "when": None,
        "day": None,
        "flag": False,
        "code": "8",
    },
    {
        "id": 3,
        "price": 0.25,
        "name": "Bob",
        "when": datetime.datetime(2020, 1, 2),
        "day": datetime.date(2021, 5, 7),
        "flag": None,
        "code": "9",
    },
]


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "data.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "data"
    sheet.append([])  # a blank row before the header
    sheet.append(["id", "price", "name", "when", "day", "flag", "code"])
    for row in EXPECTED:
        values = dict(row)
        if values["day"] is not None:
            values["day"] = datetime.datetime.combine(values["day"], datetime.time())
        sheet.append(list(values.values()))
    sheet["G5"] = 9  # a number in a text column
    second = book.create_sheet("second")
    second.append(["a", "b"])
    second.append([1, "x"])
    second.append([2, "y"])
    book.save(path)
    return str(path)


def _cli(*args):
    result = runner.invoke(app, list(args))
    assert result.exit_code == 0, result.output
    return result.output


def test_typed_values(workbook):
    assert excel_table(workbook).to_pylist() == EXPECTED


def test_column_types(workbook):
    types = {field.name: str(field.type) for field in excel_table(workbook).schema}
    assert types == {
        "id": "int64",
        "price": "double",
        "name": "string",
        "when": "timestamp[ms]",
        "day": "date32[day]",
        "flag": "bool",
        "code": "string",
    }


def test_sheets(workbook):
    assert sheet_names(workbook) == ["data", "second"]
    assert excel_table(workbook, {"table": "second"}).num_rows == 2
    assert excel_table(workbook, {"page": 1}).column_names == ["a", "b"]


def test_python_reader(workbook):
    reader = open_iterable(workbook, mode="r", iterableargs={})
    try:
        assert reader.has_totals() and reader.totals() == 3
        assert reader.list_tables() == ["data", "second"]
        assert reader.is_flat()
        assert list(reader) == EXPECTED
        reader.reset()
        assert reader.read_bulk(2) == EXPECTED[:2]
    finally:
        reader.close()


def test_row_source_engines_agree(workbook):
    source = RowSource(workbook)
    assert source.duckdb_readable()
    assert source.arrow_reader().read_all().to_pylist() == list(source) == EXPECTED


@pytest.mark.parametrize(
    "command",
    [
        ["head", "{in}", "-O", "jsonl"],
        ["count", "{in}"],
        ["select", "-f", "id,name", "{in}", "-O", "jsonl"],
        ["sort", "--by", "price", "--desc", "{in}", "-O", "jsonl"],
        ["frequency", "-f", "flag", "{in}"],
        ["uniq", "-f", "code", "{in}"],
        ["count", "{in}", "--sheet", "second"],
        ["head", "{in}", "--start-page", "1", "-O", "jsonl"],
    ],
    ids=lambda command: "-".join(part for part in command if not part.startswith(("{", "-"))),
)
def test_engines_agree(workbook, command):
    args = [arg.replace("{in}", workbook) for arg in command]
    outputs = {engine: _cli(args[0], "-e", engine, *args[1:]) for engine in ("python", "duckdb")}
    assert outputs["duckdb"] == outputs["python"]


@pytest.mark.parametrize("suffix", ["jsonl", "csv", "parquet"])
def test_convert_engines_agree(workbook, tmp_path, suffix):
    # auto: DuckDB COPY for the columns it writes as the Python writer does (not timestamps)
    results = {}
    for engine in ("python", "auto"):
        out = tmp_path / f"{engine}.{suffix}"
        _cli("convert", "-e", engine, workbook, str(out))
        if suffix == "parquet":
            results[engine] = pq.read_table(out).to_pylist()
        elif suffix == "csv":
            with open(out, encoding="utf8", newline="") as handle:
                results[engine] = list(csv.DictReader(handle))
        else:
            results[engine] = [json.loads(line) for line in out.read_text().splitlines()]
    assert results["auto"] == results["python"]
    if suffix == "jsonl":
        first = results["python"][0]
        assert (first["id"], first["day"], first["flag"]) == (1, "2021-05-06", True)


@pytest.mark.parametrize("name", ["2cols6rows.xls", "2cols6rows.xlsx"])
def test_fixtures(name):
    rows = list(RowSource(str(FIXTURES / name)))
    assert len(rows) == 6
    assert rows[0] == {"id": 1, "name": "John"}


def test_analyze_reads_typed_sheets(workbook):
    from undatum.cmds.analyzer import analyze

    report = analyze(workbook, engine="auto", stats=False)
    assert [table.id for table in report.tables] == ["data", "second"]
    fields = {field.name: field.ftype for field in report.tables[0].fields}
    assert fields["id"] == "BIGINT" and fields["day"] == "DATE"
    assert report.tables[0].num_records == 3


def test_compressed_workbook_is_not_read_by_fastexcel(workbook, tmp_path):
    import gzip

    from undatum.io.excel import open_excel

    path = tmp_path / "data.xlsx.gz"
    path.write_bytes(gzip.compress(Path(workbook).read_bytes()))
    assert RowSource(str(path)).arrow_reader() is None
    assert open_excel(str(path), {}) is None


def test_whole_numbers_beyond_int64_stay_floats():
    import pyarrow as pa

    from undatum.io.excel import normalize_table

    table = pa.table({"big": pa.array([1e19, 2.0]), "exact": pa.array([2.0**60, None])})
    out = normalize_table(table)
    assert str(out.schema.field("big").type) == "double"
    assert out.column("exact").to_pylist() == [2**60, None]


def test_normalizing_does_not_import_pandas(workbook):
    import subprocess
    import sys

    code = (
        "import sys; from undatum.io.excel import excel_table; "
        f"excel_table({workbook!r}); print('pandas' in sys.modules)"
    )
    result = subprocess.run(
        [sys.executable, "-c", code], capture_output=True, text=True, check=True
    )
    assert result.stdout.strip() == "False"
